#!/usr/bin/env python3
"""Build deterministic target CSVs and the master tracker from the workbook.

Reads ``data/targets_raw.xlsx`` (all sheets), detects the real header row,
preserves blank-URL rows and workbook order, extracts hyperlinks / embedded
URLs, and writes:

  - ``data/targets/<tab_slug>.csv`` — one file per workbook tab
  - ``data/targets.csv`` — one row per normalized domain (+ no-domain rows)
  - ``docs/TARGET_LIST_REPORT.md`` — human-readable inventory report
  - ``docs/RECON_SUMMARY.md`` — recon facts by adapter family (when any exist)

Two optional inputs are merged into the tracker:

  - ``data/target_overrides.csv`` — hand decisions (category fixes, access
    profile, wave, exclusion reason), keyed by domain or ``<tab>:<row>``
  - ``data/recon/<domain>.json`` — facts written by ``scripts/recon_site.py``

The workbook itself is never auto-fixed; corrections live in the overrides
file so the next sheet export does not silently undo them.

Requires the local-only tooling: ``pip install -r requirements-dev.txt``.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
import pandas as pd
import yaml

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = PROJECT_ROOT / "data" / "targets_raw.xlsx"
TARGETS_DIR = PROJECT_ROOT / "data" / "targets"
TRACKER_PATH = PROJECT_ROOT / "data" / "targets.csv"
REPORT_PATH = PROJECT_ROOT / "docs" / "TARGET_LIST_REPORT.md"
RECON_SUMMARY_PATH = PROJECT_ROOT / "docs" / "RECON_SUMMARY.md"
CONFIG_PATH = PROJECT_ROOT / "config" / "sites.yaml"
OVERRIDES_PATH = PROJECT_ROOT / "data" / "target_overrides.csv"
RECON_DIR = PROJECT_ROOT / "data" / "recon"

URL_RE = re.compile(r"https?://[^\s<>\"')\]]+", re.I)
DOMAIN_RE = re.compile(r"^(?:www\.)?([a-z0-9.-]+\.[a-z]{2,})", re.I)

# Workbook "done" markers — seed status, but do NOT claim output files exist.
WORKBOOK_DONE_NOTE = "workbook_done_marker"

ACCESS_PROFILES = ("open", "cloudflare", "waf_browser_ua", "datacenter_blocked", "excluded")

# Recon fields copied verbatim into the tracker.
RECON_FIELDS = ("cdn", "crawl_delay", "sitemap_kind", "est_urls", "earliest_year")

# Seconds per article when a site has no configured delay yet: Northwestern's
# post-block pace for Cloudflare sites, else the crawl delay with a 6 s floor.
CLOUDFLARE_DELAY = 12.0
MIN_PLANNED_DELAY = 6.0

TAB_FIELDS = [
    "tab",
    "source_row",
    "university_name",
    "ranking_raw",
    "newspaper_raw",
    "url",
    "domain",
    "done",
    "notes",
]

TRACKER_FIELDS = [
    "site_key",
    "domain",
    "category_labels",
    "platform",
    "access_profile",
    "cdn",
    "crawl_delay",
    "sitemap_kind",
    "status",
    "est_urls",
    "earliest_year",
    "est_days",
    "wave",
    "exclusion_reason",
    "shared_with",
    "universities",
    "tabs",
    "source_rows",
    "sample_url",
    "notes",
    "workbook_done",
]


def slugify(name: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    return s or "sheet"


def normalize_label(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().lower()


def find_header_row(rows: list[list[object]]) -> int:
    """Return the 0-based index of the header row with University + Newspaper."""
    for idx, row in enumerate(rows):
        labels = [normalize_label(c) for c in row]
        has_uni = any(lab.startswith("university name") for lab in labels)
        has_paper = any("main school newspaper" in lab for lab in labels)
        if has_uni and has_paper:
            return idx
    raise ValueError("Could not find header row with University Name + Main School Newspaper")


def col_index(header: list[object], *predicates) -> int | None:
    for i, cell in enumerate(header):
        lab = normalize_label(cell)
        for pred in predicates:
            if pred(lab):
                return i
    return None


def extract_url(cell_value: object, hyperlink_target: str | None) -> tuple[str, str]:
    """Return (url, notes). Prefer hyperlink target; else first URL in text."""
    notes_parts: list[str] = []
    raw = "" if cell_value is None else str(cell_value).strip()
    url = ""
    if hyperlink_target:
        url = hyperlink_target.strip()
        if raw and raw != url and not raw.startswith("http"):
            notes_parts.append(f"label:{raw}")
    if not url and raw:
        match = URL_RE.search(raw)
        if match:
            url = match.group(0).rstrip(".,);]")
            if raw != url:
                notes_parts.append(f"embedded_in:{raw}")
        elif raw.upper() == "N/A":
            notes_parts.append("newspaper:N/A")
        elif raw:
            notes_parts.append(f"non_url:{raw}")
    return url, "; ".join(notes_parts)


def normalize_domain(url: str) -> str:
    if not url:
        return ""
    parsed = urlparse(url if "://" in url else f"https://{url}")
    host = (parsed.netloc or parsed.path.split("/")[0]).lower()
    if host.startswith("www."):
        host = host[4:]
    # Drop credentials/ports for dedupe key.
    host = host.split("@")[-1].split(":")[0]
    return host


def guess_platform(url: str, domain: str) -> str:
    if not url and not domain:
        return ""
    lower = f"{url} {domain}".lower()
    if "instagram.com" in lower:
        return "instagram"
    if any(
        token in lower
        for token in (
            "library.",
            "/library",
            "archives.",
            "digitalcollections",
            "catalog",
            "jstor",
        )
    ):
        return "library_archive"
    if any(
        token in lower
        for token in (
            "/news",
            "news.",
            "today.",
            "newsroom",
            "communications",
        )
    ) and not any(
        token in domain
        for token in (
            "daily",
            "chronicle",
            "crimson",
            "maroon",
            "spectator",
            "tribune",
            "herald",
            "times",
            "record",
            "orient",
            "phoenix",
            "bruin",
            "cavalier",
            "guardian",
        )
    ):
        # Heuristic only — never auto-delete; flagged in report.
        return "possible_university_newsroom"
    return "unknown"


def load_configured_sites(path: Path = CONFIG_PATH) -> dict[str, dict]:
    """``{domain: {"site_key", "platform", "delay"}}`` for sites in sites.yaml."""
    if not path.is_file():
        return {}
    config = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    default_rate = (config.get("defaults") or {}).get("rate_limit") or {}
    out: dict[str, dict] = {}
    for key, cfg in (config.get("sites") or {}).items():
        domain = normalize_domain(str((cfg or {}).get("base_url", "")))
        if not domain:
            continue
        rate = {**default_rate, **(cfg.get("rate_limit") or {})}
        lo = float(rate.get("delay_min", 1.0))
        hi = float(rate.get("delay_max", lo))
        out[domain] = {
            "site_key": key,
            "platform": str(cfg.get("platform", "")),
            "delay": (lo + hi) / 2,
        }
    return out


def load_overrides(path: Path = OVERRIDES_PATH) -> dict[str, dict]:
    """``{key: row}`` where key is a domain or ``<tab>:<source_row>``."""
    if not path.is_file():
        return {}
    with path.open(newline="", encoding="utf-8") as fh:
        rows = [r for r in csv.DictReader(fh) if (r.get("key") or "").strip()]
    out: dict[str, dict] = {}
    for row in rows:
        key = row["key"].strip()
        if key in out:
            raise ValueError(f"Duplicate key {key!r} in {path.name}")
        profile = (row.get("access_profile") or "").strip()
        if profile and profile not in ACCESS_PROFILES:
            raise ValueError(f"{path.name}: unknown access_profile {profile!r} for {key}")
        out[key] = {k: (v or "").strip() for k, v in row.items()}
    return out


def load_recon(recon_dir: Path = RECON_DIR) -> dict[str, dict]:
    if not recon_dir.is_dir():
        return {}
    out: dict[str, dict] = {}
    for path in sorted(recon_dir.glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("domain"):
            out[data["domain"]] = data
    return out


def planned_delay(configured: dict | None, recon: dict | None) -> float | None:
    """Average seconds per article the site will be scraped at."""
    if configured:
        return configured["delay"]
    if not recon:
        return None
    if recon.get("cdn") == "cloudflare":
        return CLOUDFLARE_DELAY
    try:
        crawl = float(recon.get("crawl_delay") or 0)
    except (TypeError, ValueError):
        crawl = 0.0
    return max(crawl, MIN_PLANNED_DELAY)


def est_days(est_urls: object, delay: float | None) -> str:
    """Lower bound: URLs x delay, ignoring fetch time and retries."""
    try:
        urls = int(est_urls)
    except (TypeError, ValueError):
        return ""
    if delay is None or urls <= 0:
        return ""
    return f"{urls * delay / 86400:.1f}"


def enrich(
    row: dict,
    configured: dict | None,
    recon: dict | None,
    override: dict | None,
) -> dict:
    """Apply sites.yaml, recon facts and hand overrides to one tracker row."""
    override = override or {}
    recon = recon or {}
    if configured:
        row["site_key"] = configured["site_key"]
    row["category_labels"] = override.get("category_labels") or row["tabs"]
    row["platform"] = (
        (configured or {}).get("platform") or recon.get("platform") or row["platform"]
    )
    for field in RECON_FIELDS:
        value = recon.get(field)
        row[field] = "" if value is None else str(value)
    row["access_profile"] = override.get("access_profile") or recon.get("access_profile", "")
    row["wave"] = override.get("wave", "")
    row["exclusion_reason"] = override.get("exclusion_reason", "")
    if row["exclusion_reason"]:
        row["access_profile"] = "excluded"
    row["est_days"] = est_days(row["est_urls"], planned_delay(configured, recon))
    if override.get("notes"):
        row["notes"] = "; ".join(n for n in (row["notes"], override["notes"]) if n)

    if row["exclusion_reason"]:
        row["status"] = "excluded"
    elif configured:
        row["status"] = "configured"
    elif recon and row["status"] == "pending_recon":
        row["status"] = "recon_done"
    return row


def duplicate_names(tab_rows: dict[str, list[dict]]) -> list[str]:
    """Report lines for a university name repeated within one tab."""
    lines: list[str] = []
    for tab_name, rows in tab_rows.items():
        counts = Counter(r["university_name"] for r in rows if r["university_name"])
        for name, n in counts.items():
            if n < 2:
                continue
            where = ", ".join(
                f"row {r['source_row']} -> {r['domain'] or '(no URL)'}"
                for r in rows
                if r["university_name"] == name
            )
            lines.append(f"- [{tab_name}] {name!r} appears {n} times: {where}")
    return lines


def seed_status(domain: str, workbook_done: bool) -> str:
    if workbook_done:
        return "workbook_done_no_local_output"
    if not domain:
        return "missing_url"
    return "pending_recon"


def read_workbook(path: Path) -> dict[str, list[dict]]:
    """Return {tab_name: [row_dicts]} preserving order."""
    wb = openpyxl.load_workbook(path, data_only=False)
    # Also load values via pandas for ranking text consistency.
    pd_sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
    out: dict[str, list[dict]] = {}

    for tab_name in wb.sheetnames:
        ws = wb[tab_name]
        raw_rows: list[list[object]] = []
        for r in range(1, ws.max_row + 1):
            raw_rows.append([ws.cell(r, c).value for c in range(1, ws.max_column + 1)])
        header_idx = find_header_row(raw_rows)
        header = raw_rows[header_idx]
        uni_i = col_index(header, lambda lab: lab.startswith("university name"))
        rank_i = col_index(header, lambda lab: lab.startswith("ranking"))
        paper_i = col_index(header, lambda lab: "main school newspaper" in lab)
        done_i = col_index(header, lambda lab: lab.startswith("done"))
        if uni_i is None or paper_i is None:
            raise ValueError(f"Missing required columns on tab {tab_name!r}")

        rows: list[dict] = []
        for offset, row in enumerate(raw_rows[header_idx + 1 :], start=header_idx + 2):
            uni = row[uni_i] if uni_i < len(row) else None
            # Skip trailing fully-empty rows.
            if uni is None and all(c is None or str(c).strip() == "" for c in row):
                continue
            paper_cell = ws.cell(offset, paper_i + 1)
            hyper = paper_cell.hyperlink.target if paper_cell.hyperlink else None
            paper_val = row[paper_i] if paper_i < len(row) else None
            # Prefer openpyxl cell value; fall back to pandas cell if needed.
            if paper_val is None and tab_name in pd_sheets:
                try:
                    paper_val = pd_sheets[tab_name].iloc[offset - 1, paper_i]
                except Exception:
                    paper_val = None
            url, notes = extract_url(paper_val, hyper)
            domain = normalize_domain(url)
            done_val = ""
            if done_i is not None and done_i < len(row) and row[done_i] is not None:
                done_val = str(row[done_i]).strip()
            ranking = ""
            if rank_i is not None and rank_i < len(row) and row[rank_i] is not None:
                ranking = str(row[rank_i]).strip()
            uni_name = "" if uni is None else str(uni).strip()
            rows.append(
                {
                    "tab": tab_name,
                    "source_row": offset,
                    "university_name": uni_name,
                    "ranking_raw": ranking,
                    "newspaper_raw": "" if paper_val is None else str(paper_val).strip(),
                    "url": url,
                    "domain": domain,
                    "done": done_val,
                    "notes": notes,
                }
            )
        out[tab_name] = rows
    return out


def write_tab_csvs(tab_rows: dict[str, list[dict]]) -> list[Path]:
    TARGETS_DIR.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for tab_name, rows in tab_rows.items():
        path = TARGETS_DIR / f"{slugify(tab_name)}.csv"
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=TAB_FIELDS)
            writer.writeheader()
            for row in rows:
                writer.writerow({k: row.get(k, "") for k in TAB_FIELDS})
        paths.append(path)
    return paths


def build_tracker(
    tab_rows: dict[str, list[dict]],
    configured: dict[str, dict] | None = None,
    recon: dict[str, dict] | None = None,
    overrides: dict[str, dict] | None = None,
) -> list[dict]:
    """One row per normalized domain; retain each no-domain row separately."""
    by_domain: dict[str, list[dict]] = defaultdict(list)
    no_domain: list[dict] = []
    for rows in tab_rows.values():
        for row in rows:
            if row["domain"]:
                by_domain[row["domain"]].append(row)
            else:
                no_domain.append(row)

    tracker: list[dict] = []
    # Preserve first-seen workbook order for domains.
    seen_domains: list[str] = []
    for rows in tab_rows.values():
        for row in rows:
            d = row["domain"]
            if d and d not in seen_domains:
                seen_domains.append(d)

    for domain in seen_domains:
        group = by_domain[domain]
        universities = [r["university_name"] for r in group if r["university_name"]]
        tabs = sorted({r["tab"] for r in group})
        source_rows = [f"{r['tab']}:{r['source_row']}" for r in group]
        workbook_done = any(str(r.get("done", "")).lower() == "done" for r in group)
        shared = universities[1:] if len(universities) > 1 else []
        notes = [r["notes"] for r in group if r.get("notes")]
        if workbook_done:
            notes.append(WORKBOOK_DONE_NOTE)
        sample_url = next((r["url"] for r in group if r["url"]), "")
        tracker.append(
            {
                "site_key": "",
                "domain": domain,
                "platform": guess_platform(sample_url, domain),
                "status": seed_status(domain, workbook_done),
                "shared_with": " | ".join(shared),
                "universities": " | ".join(universities),
                "tabs": " | ".join(tabs),
                "source_rows": " | ".join(source_rows),
                "sample_url": sample_url,
                "notes": "; ".join(notes),
                "workbook_done": "yes" if workbook_done else "",
            }
        )

    for row in no_domain:
        workbook_done = str(row.get("done", "")).lower() == "done"
        tracker.append(
            {
                "site_key": "",
                "domain": "",
                "platform": guess_platform(row.get("url", ""), ""),
                "status": seed_status("", workbook_done),
                "shared_with": "",
                "universities": row.get("university_name", ""),
                "tabs": row.get("tab", ""),
                "source_rows": f"{row.get('tab')}:{row.get('source_row')}",
                "sample_url": row.get("url", ""),
                "notes": row.get("notes", ""),
                "workbook_done": "yes" if workbook_done else "",
            }
        )

    configured = configured or {}
    recon = recon or {}
    overrides = overrides or {}
    for row in tracker:
        key = row["domain"] or row["source_rows"]
        enrich(row, configured.get(row["domain"]), recon.get(row["domain"]), overrides.get(key))
    unknown = sorted(set(overrides) - {r["domain"] or r["source_rows"] for r in tracker})
    if unknown:
        raise ValueError(f"target_overrides.csv keys match no tracker row: {unknown}")
    return tracker


def write_tracker(rows: list[dict]) -> Path:
    TRACKER_PATH.parent.mkdir(parents=True, exist_ok=True)
    with TRACKER_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=TRACKER_FIELDS)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row.get(k, "") for k in TRACKER_FIELDS})
    return TRACKER_PATH


def write_report(tab_rows: dict[str, list[dict]], tracker: list[dict]) -> Path:
    total_source = sum(len(v) for v in tab_rows.values())
    with_url = sum(1 for rows in tab_rows.values() for r in rows if r["url"])
    missing_url = total_source - with_url
    domains = [r for r in tracker if r["domain"]]
    no_domain = [r for r in tracker if not r["domain"]]
    shared = [r for r in domains if r["shared_with"]]
    workbook_done = [r for r in tracker if r.get("workbook_done") == "yes"]
    configured = [r for r in tracker if r["site_key"]]
    resolved = [
        r for r in tracker if r["exclusion_reason"] or r["category_labels"] != r["tabs"]
    ]

    suspicious: list[str] = []
    for r in tracker:
        if r in resolved:
            continue
        plat = r.get("platform", "")
        uni = r.get("universities", "")
        url = r.get("sample_url", "")
        domain = r.get("domain", "")
        if plat in {"instagram", "library_archive", "possible_university_newsroom"}:
            suspicious.append(f"- `{domain or '(no domain)'}` — {uni}: {url} [{plat}]")
        # Known mismatched mappings called out in the audit/plan.
        if "oudaily.com" in domain and "Oklahoma" in uni:
            suspicious.append(
                f"- `{domain}` — {uni}: possible wrong-school mapping (OU Daily vs Oklahoma City University)"
            )
        if "concordiensis.com" in domain:
            suspicious.append(
                f"- `{domain}` — {uni}: possible wrong-school mapping (Concordiensis / Union)"
            )

    # Deduplicate suspicious bullets while preserving order.
    seen: set[str] = set()
    suspicious_unique = []
    for line in suspicious:
        if line not in seen:
            seen.add(line)
            suspicious_unique.append(line)

    missing_lines = []
    for rows in tab_rows.values():
        for r in rows:
            if not r["url"]:
                missing_lines.append(
                    f"- [{r['tab']} row {r['source_row']}] {r['university_name'] or '(blank university)'} "
                    f"— raw={r['newspaper_raw']!r}"
                )

    lines = [
        "# Target List Report",
        "",
        "Generated by `scripts/build_targets.py` from `data/targets_raw.xlsx`.",
        "Suspicious rows are flagged only; nothing is auto-fixed or deleted.",
        "",
        "## Totals",
        "",
        f"- Workbook tabs: {len(tab_rows)}",
        f"- Source university rows: {total_source}",
        f"- Rows with a URL: {with_url}",
        f"- Rows missing a URL: {missing_url}",
        f"- Unique domains: {len(domains)}",
        f"- No-domain tracker rows: {len(no_domain)}",
        f"- Domains shared by multiple universities: {len(shared)}",
        f"- Workbook `done` markers: {len(workbook_done)} "
        "(seeded as `workbook_done_no_local_output` unless a built site_key exists)",
        f"- Sites configured in `config/sites.yaml`: {len(configured)} "
        f"({', '.join(r['site_key'] for r in configured)})",
        "",
        "## Per-tab counts",
        "",
        "| Tab | Rows | With URL | Missing URL |",
        "| --- | ---: | ---: | ---: |",
    ]
    for tab_name, rows in tab_rows.items():
        wu = sum(1 for r in rows if r["url"])
        lines.append(f"| {tab_name} | {len(rows)} | {wu} | {len(rows) - wu} |")

    lines.extend(
        [
            "",
            "## Progress by category",
            "",
            "Counts use `category_labels` (sheet tabs after overrides); a paper",
            "listed in two categories counts in both.",
            "",
            "| Category | Domains | Excluded | Recon done | Configured | In a wave |",
            "| --- | ---: | ---: | ---: | ---: | ---: |",
        ]
    )
    for tab_name in tab_rows:
        in_cat = [
            r for r in domains if tab_name in r["category_labels"].split(" | ")
        ]
        lines.append(
            f"| {tab_name} | {len(in_cat)} "
            f"| {sum(1 for r in in_cat if r['exclusion_reason'])} "
            f"| {sum(1 for r in in_cat if r['cdn'] or r['sitemap_kind'])} "
            f"| {sum(1 for r in in_cat if r['site_key'])} "
            f"| {sum(1 for r in in_cat if r['wave'])} |"
        )

    lines.extend(["", "## Resolved by `data/target_overrides.csv`", ""])
    if not resolved:
        lines.append("_None._")
    for r in resolved:
        label = r["domain"] or r["source_rows"]
        if r["exclusion_reason"]:
            lines.append(f"- `{label}` — {r['universities']}: excluded ({r['exclusion_reason']})")
        else:
            lines.append(
                f"- `{label}` — {r['universities']}: categories "
                f"{r['tabs']!r} -> {r['category_labels']!r}"
            )

    dupes = duplicate_names(tab_rows)
    lines.extend(["", "## Repeated university names within a tab", ""])
    lines.extend(dupes or ["_None._"])

    lines.extend(
        [
            "",
            "## Shared newspapers (same domain, multiple universities)",
            "",
        ]
    )
    if not shared:
        lines.append("_None detected._")
    else:
        for r in shared:
            lines.append(
                f"- `{r['domain']}` — {r['universities']} (tabs: {r['tabs']})"
            )

    lines.extend(["", "## Missing URLs", ""])
    if not missing_lines:
        lines.append("_None._")
    else:
        lines.extend(missing_lines)

    lines.extend(
        [
            "",
            "## Reachability caveats",
            "",
            "- A prior one-request-per-domain probe is **not** re-run here.",
            "- HTTP 403/401/429 on a HEAD/GET probe does **not** mean the site is dead;",
            "  Duke and Yale already require browser UA or Playwright for access.",
            "- Instagram-only, library/catalog, and university news-office URLs need",
            "  human review before recon or builder work.",
            "",
            "## Suspicious / review mappings",
            "",
        ]
    )
    if not suspicious_unique:
        lines.append("_None auto-flagged._")
    else:
        lines.extend(suspicious_unique)

    lines.extend(
        [
            "",
            "## Configured sites vs workbook `done`",
            "",
            "- Sites with a `config/sites.yaml` entry have status `configured`;",
            "  that says nothing about how much output exists.",
            "- Workbook marks MIT (`thetech.com`) and Columbia (`columbiaspectator.com`) as",
            "  `done`, but **no MIT or Columbia output CSVs** are in this repository.",
            "  Tracker status for those domains is `workbook_done_no_local_output`.",
            "",
        ]
    )

    REPORT_PATH.parent.mkdir(parents=True, exist_ok=True)
    REPORT_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return REPORT_PATH


def write_recon_summary(tracker: list[dict], recon: dict[str, dict]) -> Path:
    """One table row per reconned domain, grouped by suggested adapter family."""
    rows = [r for r in tracker if r["domain"] in recon]
    fill_cols = ("publication_date", "author", "text", "section", "title")
    lines = [
        "# Recon summary",
        "",
        "Generated by `scripts/build_targets.py` from `data/recon/*.json`",
        "(written by `scripts/recon_site.py`, honest research UA, from the Mac's IP).",
        "Access from the server still needs `scripts/access_probe.py`.",
        "`est_days` is a lower bound: estimated URLs x planned delay.",
        "Field fill is out of 3 sampled articles (oldest, middle, newest).",
        "",
        f"Domains reconned: {len(rows)}",
    ]
    for family in ("sno", "wordpress", "snworks", "blox", "unknown"):
        group = [r for r in rows if (r["platform"] or "unknown") == family]
        if not group:
            continue
        lines.extend(
            [
                "",
                f"## {family} ({len(group)})",
                "",
                "| Domain | Categories | Access | CDN | Home | Crawl delay | Sitemap | Est. URLs | Est. days "
                "| Earliest | Date | Author | Body | Section | Title |",
                "| --- | --- | --- | --- | ---: | ---: | --- | ---: | ---: | ---: | --- | --- | --- | --- | --- |",
            ]
        )
        for r in sorted(group, key=lambda r: (r["category_labels"], r["domain"])):
            data = recon[r["domain"]]
            fill = data.get("field_fill") or {}
            lines.append(
                f"| {r['domain']} | {r['category_labels']} | {r['access_profile']} | {r['cdn']} "
                f"| {data.get('home_status') or ''} | {r['crawl_delay']} | {r['sitemap_kind']} "
                f"| {r['est_urls']} | {r['est_days']} | {r['earliest_year']} | "
                + " | ".join(fill.get(c, "") for c in fill_cols)
                + " |"
            )
    RECON_SUMMARY_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return RECON_SUMMARY_PATH


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--xlsx",
        type=Path,
        default=DEFAULT_XLSX,
        help="Path to targets_raw.xlsx",
    )
    args = parser.parse_args(argv)
    tab_rows = read_workbook(args.xlsx)
    tab_paths = write_tab_csvs(tab_rows)
    recon = load_recon()
    tracker = build_tracker(tab_rows, load_configured_sites(), recon, load_overrides())
    tracker_path = write_tracker(tracker)
    report_path = write_report(tab_rows, tracker)
    print(f"Wrote {len(tab_paths)} tab CSV(s) under {TARGETS_DIR}")
    print(f"Wrote tracker -> {tracker_path} ({len(tracker)} rows)")
    print(f"Wrote report  -> {report_path}")
    if recon:
        print(f"Wrote recon   -> {write_recon_summary(tracker, recon)} ({len(recon)} domains)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
